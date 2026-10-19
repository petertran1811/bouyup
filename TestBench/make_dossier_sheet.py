#!/usr/bin/env python3
"""results/dossiers_round*.jsonl → results/TestBench_50_ho_so.xlsx (the 50-dossier bypass bench).

Round 1 = before the fixes of 17/10/2026, round 2 = after. Run the bench first:
    ../BackEnd/.venv/bin/python -m pytest cases/06_dossiers_bypass ; cp results/dossiers.jsonl results/dossiers_round2.jsonl
"""

from __future__ import annotations

import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
RES = HERE / "results"
sys.path.insert(0, str(HERE / "cases" / "06_dossiers_bypass"))
sys.path.insert(0, str(HERE.parent / "BackEnd" / "src"))

PLUM, ROSE, PEACH, INK = "7D4A78", "F7EBF1", "FCEFE6", "2B1D26"
OK_FILL, BAD_FILL = PatternFill("solid", fgColor="E8F3EC"), PatternFill("solid", fgColor="FBE8E6")
HEAD = PatternFill("solid", fgColor=PLUM)
THIN = Side(style="thin", color="E2D4DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

FIXES = {
    "DS-07": ("Bẫy khởi động lạnh: người bán nhỏ có FATE mỏng bị cộng dồn trần FATE mỏng × hệ số khoản vay mồi 50% nên hạn "
              "mức rơi dưới mức vay tối thiểu 10.000 → không bao giờ vay được khoản đầu để có lịch sử.",
              "Khoản vay mồi không xuống dưới mức tối thiểu khi hạn mức tính được đạt mức tối thiểu.",
              "BackEnd/src/app/core/algorithm.py · limit_caps"),
    "DS-29": ("Khai lố doanh thu tự khai trong ngưỡng chịu đựng 15% (ví dụ +12%) vẫn được tính là “khớp thuế” và cộng "
              "+0,05 hệ số thận trọng → hạn mức cao hơn mà không cần bằng chứng.",
              "Số doanh thu tự khai chỉ được cộng hệ số khi có ảnh chụp chứa đúng các số đó (OCR đối chiếu).",
              "BackEnd/src/app/agents/agent1_verify.py · tham số revenue.bonus_requires_images"),
    "DS-40": ("Form không còn ô nhãn hiệu nên hàng nhái thương hiệu (“Giày Nike replica 1:1”) đi qua kiểm tra sở hữu trí "
              "tuệ và được duyệt tự động.",
              "Đọc tên hàng và mô tả “hàng khác”: từ khóa hàng nhái → coi là không có ủy quyền (LAW-IP-01, chuyển cán bộ); "
              "thương hiệu lớn không khai ủy quyền → “chưa rõ” (LAW-IP-02).",
              "BackEnd/src/app/data_sources/sources.py · ip_signals; agent3_control.evaluate_rules"),
}
RESIDUAL = [
    ("Nhận diện hàng cấm / hàng nhái dựa trên từ khóa", "Người bán có thể viết lách (ví dụ “p0d”, ảnh thay chữ). Nên bổ sung "
     "LLM phân loại mặt hàng (đã có khung cho Agent 3) và đối chiếu danh mục hàng của cửa hàng trên sàn."),
    ("Ảnh doanh thu chỉ được đối chiếu số, không kiểm chỉnh sửa ảnh", "Ảnh tự vẽ khớp số tự khai vẫn bị chặn nhờ đối chiếu "
     "thuế (DS-31), nhưng chưa có kiểm tra pháp y ảnh (EXIF, dấu vết chỉnh sửa)."),
    ("CCCD: OCR không xác minh thẻ thật", "OCR chỉ đọc trường và kiểm checksum, hạn, tuổi; không đọc chip NFC, không tra cơ "
     "quan công an. Thẻ in giả với số hợp lệ vẫn qua bước đọc, khuôn mặt (ArcFace 1:1) là lớp chặn tiếp theo."),
    ("Bỏ ô nhà cung cấp", "Hệ thống dùng hóa đơn mua vào của chính mặt hàng thay cho nhà cung cấp khai; hàng mới chưa có "
     "hóa đơn bị tính hệ số 0,7. Việc kiểm dùng vốn sau giải ngân chấp nhận bất kỳ nhà cung cấp có lịch sử, không liên kết."),
    ("19 ca TestBench cũ (nhóm 01–05) không chạy được", "Viết cho mô hình ủy quyền và API cũ (trước loan circle); fail y "
     "hệt trên bản code trước 17/10. Cần viết lại theo luồng hiện tại."),
]


def load(name: str) -> dict[str, dict]:
    path = RES / name
    if not path.exists():
        return {}
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return {r["id"]: r for r in rows}


def head(ws, row: int, values: list[str]) -> None:
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), HEAD, WRAP, BOX


def widths(ws, cols: list[int]) -> None:
    for i, w in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def outcome(r: dict | None) -> str:
    if not r:
        return "—"
    d = r.get("detail") or {}
    parts = []
    if d.get("status") is not None:
        parts.append(f"trạng thái {d['status']}")
    if d.get("limit"):
        parts.append(f"hạn mức {d['limit']:,}".replace(",", "."))
    if d.get("flags"):
        parts.append("cờ " + ", ".join(d["flags"]))
    if d.get("reason"):
        parts.append(d["reason"])
    for k in ("prudence", "select", "second_sign", "loans", "decision", "reauthorize", "sign", "junk", "duplicate",
              "app", "image", "name_check", "card", "hidden_from_seller", "alerts", "images", "b3", "tenors"):
        if k in d and d[k] not in (None, [], {}):
            parts.append(f"{k}: {d[k]}")
    if r.get("error"):
        parts.append("lỗi luồng: " + r["error"][:120])
    return " · ".join(str(p) for p in parts)[:600]


def main() -> None:
    from dossiers import ARCHETYPES, VND_PER_CNY  # noqa: E402

    r1, r2 = load("dossiers_round1.jsonl"), load("dossiers_round2.jsonl") or load("dossiers.jsonl")
    ids = sorted(r2 or r1)
    rows = [(i, r1.get(i), r2.get(i)) for i in ids]
    wb = Workbook()

    # --- Tổng quan
    ws = wb.active
    ws.title = "Tổng quan"
    ws["A1"] = "TestBench · 50 hồ sơ vay BuoyUp (17/10/2026)"
    ws["A1"].font = Font(bold=True, size=16, color=PLUM)
    ws["A2"] = ("Mỗi hồ sơ chạy trọn luồng thật qua API: quét CCCD (OCR) → khuôn mặt → cho phép kiểm tra → kế hoạch có "
                "cam đoan → xác minh → hạn mức → ký → giải ngân. ĐẠT = hệ thống dừng đúng chỗ (phục vụ đúng người trung "
                "thực, chặn đúng người gian). Lần 1 = trước khi vá, lần 2 = sau khi vá.")
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 60
    head(ws, 4, ["Nhóm", "Số hồ sơ", "Lần 1 đạt", "Lần 2 đạt", "Tỷ lệ lần 2", "Ghi chú"])
    groups = [("Tất cả", lambda d: True)] + [(f"Độ khó {n}", lambda d, n=n: d["difficulty"] == n) for n in range(1, 6)] + [
        (k.capitalize(), lambda d, k=k: d["kind"] == k) for k in ("trung thực", "khai gian", "lạm dụng luồng")]
    line = 5
    for label, pred in groups:
        sel = [(a, b) for _i, a, b in rows if pred(b or a)]
        if not sel:
            continue
        p1 = sum(1 for a, _b in sel if a and a["secure"])
        p2 = sum(1 for _a, b in sel if b and b["secure"])
        vals = [label, len(sel), p1 if r1 else "—", p2, f"{p2 / len(sel):.0%}", ""]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=line, column=col, value=v)
            c.border = BOX
            if label == "Tất cả":
                c.font = Font(bold=True)
        line += 1
    line += 1
    ws.cell(row=line, column=1, value="Lỗ hổng tìm thấy ở lần 1 và cách vá").font = Font(bold=True, size=13, color=PLUM)
    line += 1
    head(ws, line, ["Hồ sơ", "Lỗ hổng", "Cách vá", "Vị trí", "Lần 1", "Lần 2"])
    for i, (bug, fix, where) in FIXES.items():
        line += 1
        vals = [i, bug, fix, where, "KHÔNG ĐẠT" if r1.get(i) and not r1[i]["secure"] else "đạt",
                "ĐẠT" if r2.get(i, {}).get("secure") else "KHÔNG ĐẠT"]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=line, column=col, value=v)
            c.alignment, c.border = WRAP, BOX
        ws.row_dimensions[line].height = 75
    line += 2
    ws.cell(row=line, column=1, value="Rủi ro còn lại (chưa có ca kiểm thử chặn được hoàn toàn)").font = Font(bold=True, size=13, color=PLUM)
    line += 1
    head(ws, line, ["Rủi ro", "Mô tả và đề xuất"])
    for risk, desc in RESIDUAL:
        line += 1
        for col, v in enumerate((risk, desc), 1):
            c = ws.cell(row=line, column=col, value=v)
            c.alignment, c.border = WRAP, BOX
        ws.merge_cells(start_row=line, start_column=2, end_row=line, end_column=6)
        ws.row_dimensions[line].height = 48
    widths(ws, [26, 52, 46, 34, 12, 12])

    # --- 50 hồ sơ
    ws = wb.create_sheet("50 hồ sơ")
    head(ws, 1, ["Mã", "Hồ sơ", "Mẫu gốc", "Tư cách vay", "Độ khó", "Loại", "Cách lách", "Kết quả an toàn mong đợi",
                 "Lần 1", "Lần 2", "Kết quả thực tế (lần 2)", "Giây"])
    bt = {"individual": "Cá nhân", "individual_business": "Hộ kinh doanh", "company": "Doanh nghiệp"}
    for n, (i, a, b) in enumerate(rows, 2):
        d = b or a
        vals = [i, d["title"], d["archetype"], bt.get(d["borrower_type"], d["borrower_type"]), d["difficulty"], d["kind"],
                d["technique"], d["expect"], ("ĐẠT" if a["secure"] else "KHÔNG ĐẠT") if a else "—",
                ("ĐẠT" if b["secure"] else "KHÔNG ĐẠT") if b else "—", outcome(b), (b or {}).get("seconds")]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=n, column=col, value=v)
            c.alignment, c.border = WRAP, BOX
        for col, res in ((9, a), (10, b)):
            if res:
                ws.cell(row=n, column=col).fill = OK_FILL if res["secure"] else BAD_FILL
                ws.cell(row=n, column=col).font = Font(bold=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L{len(rows) + 1}"
    widths(ws, [8, 40, 26, 14, 8, 14, 30, 38, 12, 12, 60, 7])

    # --- Hồ sơ gốc
    ws = wb.create_sheet("Hồ sơ gốc")
    head(ws, 1, ["Mẫu", "Doanh thu tháng (VND)", "Doanh thu tháng (tệ)", "Điểm cửa hàng", "Biên lãi gộp", "Ngành",
                 "Mặt hàng kế hoạch", "Số lượng", "Giá nhập (VND)", "Giá nhập (tệ)"])
    for n, a in enumerate(ARCHETYPES.values(), 2):
        vals = [a["label"], a["monthly_vnd"], round(a["monthly_vnd"] / VND_PER_CNY), a["rating"], a["margin"],
                a["industry"], a["item"], a["qty"], a["price_vnd"], round(a["price_vnd"] / VND_PER_CNY)]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=n, column=col, value=v)
            c.border = BOX
            if col in (2, 3, 9, 10):
                c.number_format = "#,##0"
    ws.cell(row=7, column=1, value=(f"Nguồn: Backend_CCCD/chuan_ho_so (4 hồ sơ chuẩn của nhóm, ngày chốt 01/10/2026); "
                                    f"quy đổi {VND_PER_CNY:,} VND = 1 tệ. Mỗi hồ sơ trong bench là một biến thể "
                                    "(quy mô, tư cách vay, tín dụng, cách khai) dựng trên 4 mẫu này.").replace(",", "."))
    ws.merge_cells("A7:J7")
    ws["A7"].alignment = WRAP
    ws.row_dimensions[7].height = 40
    widths(ws, [30, 20, 18, 14, 12, 12, 26, 10, 16, 14])

    out = RES / "TestBench_50_ho_so.xlsx"
    wb.save(out)
    print(f"wrote {out} · round 1 {sum(1 for _i, a, _b in rows if a and a['secure'])}/{len(rows)} · "
          f"round 2 {sum(1 for _i, _a, b in rows if b and b['secure'])}/{len(rows)}")


if __name__ == "__main__":
    main()
